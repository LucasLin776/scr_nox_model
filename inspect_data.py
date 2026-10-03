import openpyxl, json, os, numpy as np, pandas as pd
p='/Users/shuoqinlin/靖海电厂1号600MW机组脱硝相关数据（0MW-600MW).xlsx'
wb=openpyxl.load_workbook(p,read_only=True,data_only=True)
ws=wb[wb.sheetnames[0]]
rows=ws.iter_rows(values_only=True); raw_headers=list(next(rows))
def uniq(headers):
 out=[]; seen={}
 for i,h in enumerate(headers):
  base=str(h).replace('\\t',' ').replace('\t',' ').strip()
  if base in seen: seen[base]+=1; base=f'{base}__dup{seen[base]}'
  else: seen[base]=1
  out.append(base)
 return out
headers=uniq(raw_headers)
data=list(rows)
df=pd.DataFrame(data,columns=headers)
print('sheets',wb.sheetnames,'nrows',len(df),'ncols',len(df.columns))
for i,(raw,h) in enumerate(zip(raw_headers,headers)): print(i,h)
print('time range',df.iloc[:,0].min(),df.iloc[:,0].max(),'seconds',df.iloc[:,1].min(),df.iloc[:,1].max(),'unique sec',df.iloc[:,1].nunique())
summary=[]
for i,c in enumerate(df.columns):
 s=pd.to_numeric(df.iloc[:,i],errors='coerce')
 summary.append({'index':i,'col':c,'missing':int(s.isna().sum()),'nunique':int(s.nunique(dropna=True)),'min':float(s.min()) if s.notna().any() else None,'max':float(s.max()) if s.notna().any() else None,'mean':float(s.mean()) if s.notna().any() else None,'std':float(s.std()) if s.notna().any() else None,'q01':float(s.quantile(.01)) if s.notna().any() else None,'q99':float(s.quantile(.99)) if s.notna().any() else None})
print(json.dumps(summary,ensure_ascii=False,indent=2))
# correlations of numeric vars
num=df.iloc[:,1:].apply(pd.to_numeric,errors='coerce')
cor=num.corr()
cor.to_csv('/Users/shuoqinlin/scr_nox_model/correlation.csv')
# high correlations with key targets
for target_idx in [2,5,10]:
 print('\nTARGET',target_idx,headers[target_idx])
 print(cor.iloc[:,target_idx].sort_values(key=lambda s:abs(s),ascending=False).head(12).to_string())
# save data and metadata
# datetime as string for portability
df.to_csv('/Users/shuoqinlin/scr_nox_model/raw_data.csv',index=False)
with open('/Users/shuoqinlin/scr_nox_model/summary.json','w') as f: json.dump(summary,f,ensure_ascii=False,indent=2)
with open('/Users/shuoqinlin/scr_nox_model/columns.json','w') as f: json.dump({'raw':raw_headers,'unique':headers},f,ensure_ascii=False,indent=2)
